import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIG
# ============================================================

DATA_DIR = Path("data")

OUTPUT_FILE = Path(
    "FrontierAtlas_Final.xlsx"
)

FILES = {
    "Startups":
        DATA_DIR / "startups_1000.json",

    "Products":
        DATA_DIR / "products_1000.json",

    "Research Papers":
        DATA_DIR / "research_papers_1000_enriched.json",

    "Jobs":
        DATA_DIR / "jobs.json",

    "News":
        DATA_DIR / "news.json",
}


# ============================================================
# JSON LOADER
# ============================================================

def load_json(path):

    if not path.exists():

        raise FileNotFoundError(
            f"File not found: {path}"
        )

    with open(
        path,
        "r",
        encoding="utf-8"
    ) as file:

        data = json.load(file)

    if not isinstance(data, list):

        raise ValueError(
            f"{path} must contain a JSON list."
        )

    return data


# ============================================================
# NESTED VALUE HELPERS
# ============================================================

def get_value(
    record,
    *keys
):

    value = record

    for key in keys:

        if not isinstance(
            value,
            dict
        ):
            return None

        value = value.get(
            key
        )

    return value


def first_value(
    record,
    paths
):

    for path in paths:

        value = get_value(
            record,
            *path
        )

        if value is not None and value != "":

            return value

    return None


def stringify(value):

    if value is None:
        return ""

    if isinstance(
        value,
        (list, dict)
    ):

        return json.dumps(
            value,
            ensure_ascii=False
        )

    return value


# ============================================================
# STARTUPS
# ============================================================

def build_startups(records):

    headers = [
        "name",
        "description",
        "website",
        "location",
        "industry",
        "funding",
        "source_url",
        "source",
    ]

    rows = [
        headers
    ]

    for record in records:

        rows.append(
            [
                stringify(
                    first_value(
                        record,
                        [
                            ("name",),
                            ("content", "name"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("description",),
                            ("content", "description"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("website",),
                            ("content", "website"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("location",),
                            ("content", "location"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("industry",),
                            ("category",),
                            ("content", "industry"),
                            ("content", "category"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("funding",),
                            ("content", "funding"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("source_url",),
                            ("source", "url"),
                            ("content", "source_url"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("source", "name"),
                            ("source",),
                        ]
                    )
                ),
            ]
        )

    return rows


# ============================================================
# PRODUCTS
# ============================================================

def build_products(records):

    headers = [
        "name",
        "description",
        "rating",
        "source_url",
        "source",
        "category_url",
    ]

    rows = [
        headers
    ]

    for record in records:

        rows.append(
            [
                stringify(
                    first_value(
                        record,
                        [
                            ("name",),
                            ("content", "name"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("description",),
                            ("content", "description"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("rating",),
                            ("content", "rating"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("source_url",),
                            ("content", "url"),
                            ("source", "url"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("source", "name"),
                            ("source",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("category_url",),
                            ("content", "category_url"),
                        ]
                    )
                ),
            ]
        )

    return rows


# ============================================================
# RESEARCH PAPERS
# ============================================================

def build_research_papers(records):

    headers = [
        "title",
        "authors",
        "paper_url",
        "published_date",
        "github_url",
        "github_stars",
        "source",
    ]

    rows = [
        headers
    ]

    for record in records:

        rows.append(
            [
                stringify(
                    first_value(
                        record,
                        [
                            ("title",),
                            ("content", "title"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("authors",),
                            ("content", "authors"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("paper_url",),
                            ("content", "paper_url"),
                            ("source", "url"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("published_date",),
                            ("content", "published_date"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("github_url",),
                            ("content", "github_url"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("github_stars",),
                            ("content", "github_stars"),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("source", "name"),
                            ("source",),
                        ]
                    )
                ),
            ]
        )

    return rows


# ============================================================
# JOBS
# ============================================================

def build_jobs(records):

    headers = [
        "title",
        "company",
        "description",
        "location",
        "category",
        "url",
        "published_at",
        "source",
    ]

    rows = [
        headers
    ]

    for record in records:

        rows.append(
            [
                stringify(
                    first_value(
                        record,
                        [
                            ("content", "title"),
                            ("title",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("content", "company"),
                            ("company",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("content", "description"),
                            ("description",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("content", "location"),
                            ("location",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("content", "category"),
                            ("category",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("content", "url"),
                            ("url",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("content", "published_at"),
                            ("published_at",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("source", "name"),
                            ("source",),
                        ]
                    )
                ),
            ]
        )

    return rows


# ============================================================
# NEWS
# ============================================================

def build_news(records):

    headers = [
        "title",
        "description",
        "url",
        "author",
        "published_at",
        "source",
    ]

    rows = [
        headers
    ]

    for record in records:

        rows.append(
            [
                stringify(
                    first_value(
                        record,
                        [
                            ("content", "title"),
                            ("title",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("content", "description"),
                            ("description",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("content", "url"),
                            ("url",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("content", "author"),
                            ("author",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("content", "published_at"),
                            ("published_at",),
                        ]
                    )
                ),

                stringify(
                    first_value(
                        record,
                        [
                            ("source", "name"),
                            ("source",),
                        ]
                    )
                ),
            ]
        )

    return rows


# ============================================================
# ENTITY MAPPING LOG
# ============================================================

def build_entity_mapping_log():

    headers = [
        "raw_name",
        "canonical_name",
        "match_method",
        "confidence",
        "source_url",
    ]

    rows = [
        headers
    ]

    # Search all JSON files for ACTUAL mapping records.
    #
    # We do not invent mappings.

    for path in DATA_DIR.glob(
        "*.json"
    ):

        try:

            data = load_json(
                path
            )

        except Exception:

            continue

        for record in data:

            if not isinstance(
                record,
                dict
            ):
                continue

            mapping = record.get(
                "mapping",
                record
            )

            if not isinstance(
                mapping,
                dict
            ):
                continue

            if "raw_name" not in mapping:

                continue

            if (
                "canonical_name"
                not in mapping
                and
                "match_method"
                not in mapping
            ):

                continue

            rows.append(
                [
                    stringify(
                        mapping.get(
                            "raw_name"
                        )
                    ),

                    stringify(
                        mapping.get(
                            "canonical_name"
                        )
                    ),

                    stringify(
                        mapping.get(
                            "match_method"
                        )
                    ),

                    stringify(
                        mapping.get(
                            "confidence"
                        )
                    ),

                    stringify(
                        mapping.get(
                            "source_url"
                        )
                    ),
                ]
            )

    return rows


# ============================================================
# EXCEL WRITER
# ============================================================

def write_sheet(
    workbook,
    sheet_name,
    rows
):

    worksheet = workbook.create_sheet(
        title=sheet_name
    )

    for row in rows:

        worksheet.append(
            row
        )

    # Freeze header row.
    worksheet.freeze_panes = "A2"

    # Add filters.
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    # Reasonable column widths.
    for column_cells in worksheet.columns:

        column_index = (
            column_cells[0].column
        )

        max_length = 0

        for cell in column_cells[:100]:

            if cell.value is None:
                continue

            length = len(
                str(cell.value)
            )

            if length > max_length:
                max_length = length

        width = min(
            max(
                max_length + 2,
                12
            ),
            50
        )

        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = width


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 70)
    print(
        "FRONTIER ATLAS — EXCEL EXPORT"
    )
    print("=" * 70)

    # --------------------------------------------------------
    # Load datasets
    # --------------------------------------------------------

    startups = load_json(
        FILES["Startups"]
    )

    products = load_json(
        FILES["Products"]
    )

    papers = load_json(
        FILES["Research Papers"]
    )

    jobs = load_json(
        FILES["Jobs"]
    )

    news = load_json(
        FILES["News"]
    )

    # --------------------------------------------------------
    # Validate required counts
    # --------------------------------------------------------

    print("\nDATA VALIDATION")

    print(
        f"Startups: "
        f"{len(startups)}"
    )

    print(
        f"Products: "
        f"{len(products)}"
    )

    print(
        f"Research Papers: "
        f"{len(papers)}"
    )

    print(
        f"Jobs: "
        f"{len(jobs)}"
    )

    print(
        f"News: "
        f"{len(news)}"
    )

    if len(startups) < 1000:

        raise ValueError(
            "Startups has fewer than 1000 records."
        )

    if len(products) < 1000:

        raise ValueError(
            "Products has fewer than 1000 records."
        )

    if len(papers) < 1000:

        raise ValueError(
            "Research Papers has fewer than 1000 records."
        )

    # --------------------------------------------------------
    # Build workbook
    # --------------------------------------------------------

    workbook = Workbook()

    # Remove default Sheet.
    default_sheet = workbook.active

    workbook.remove(
        default_sheet
    )

    # --------------------------------------------------------
    # Build six required tabs
    # --------------------------------------------------------

    print(
        "\nBUILDING WORKBOOK"
    )

    sheets = {
        "Startups":
            build_startups(
                startups
            ),

        "Products":
            build_products(
                products
            ),

        "Research Papers":
            build_research_papers(
                papers
            ),

        "Jobs":
            build_jobs(
                jobs
            ),

        "News":
            build_news(
                news
            ),

        "Entity Mapping Log":
            build_entity_mapping_log(),
    }

    # --------------------------------------------------------
    # Write
    # --------------------------------------------------------

    for sheet_name, rows in sheets.items():

        write_sheet(
            workbook,
            sheet_name,
            rows
        )

        print(
            f"{sheet_name}: "
            f"{len(rows) - 1} records"
        )

    # --------------------------------------------------------
    # Final sheet-order validation
    # --------------------------------------------------------

    expected = [
        "Startups",
        "Products",
        "Research Papers",
        "Jobs",
        "News",
        "Entity Mapping Log",
    ]

    actual = workbook.sheetnames

    if actual != expected:

        raise RuntimeError(
            "Incorrect sheet structure.\n"
            f"Expected: {expected}\n"
            f"Actual: {actual}"
        )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    workbook.save(
        OUTPUT_FILE
    )

    print("\n" + "=" * 70)
    print(
        "EXCEL EXPORT COMPLETE"
    )
    print("=" * 70)

    print(
        f"File: "
        f"{OUTPUT_FILE.resolve()}"
    )

    print(
        "\nSheets:"
    )

    for sheet_name in actual:

        print(
            f"  ✓ {sheet_name}"
        )


if __name__ == "__main__":
    main()
